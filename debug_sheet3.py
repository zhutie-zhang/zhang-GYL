import sys, os, re, glob, shutil
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
for key in list(sys.modules.keys()):
    if 'pricing' in key: del sys.modules[key]
pycache = r'C:\Users\1\Documents\Default Project\__pycache__'
if os.path.exists(pycache): shutil.rmtree(pycache)

import openpyxl
files = glob.glob(r'E:\新建文件夹\海外仓\正式价格\*.xlsx')
files = [f for f in files if '~$' not in f]
lc = [f for f in files if '乐舱' in os.path.basename(f)]
wb = openpyxl.load_workbook(lc[0], read_only=True, data_only=True)

for sn in ['Sheet3', 'SVIP FEDEX']:
    if sn not in wb.sheetnames:
        continue
    ws = wb[sn]
    rows = list(ws.iter_rows(values_only=True))
    print('=== %s zone detection ===' % sn)
    zone_cols = {}
    for ri, row in enumerate(rows[:15]):
        for ci, c in enumerate(row):
            if c is not None:
                cs = str(c).strip()
                m = re.match(r'(?:zone\s*)?(\d+)$', cs, re.I)
                if m:
                    zone_cols[ci] = 'Zone' + m.group(1)
                    print('  R%d C%d: %s -> Zone%s' % (ri, ci, cs, m.group(1)))
    print('  zone_cols count:', len(zone_cols))

wb.close()

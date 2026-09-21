import sys, os, re, glob
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
import openpyxl

files = glob.glob(r'E:\新建文件夹\海外仓\正式价格\*.xlsx')
smart = [f for f in files if 'SMART' in os.path.basename(f) and not os.path.basename(f).startswith('~$')]
leecang_new = [f for f in files if '新乐舱' in os.path.basename(f) and not os.path.basename(f).startswith('~$')]

for f in smart:
    print('=== SMART: %s ===' % os.path.basename(f))
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2: continue
        text_all = ' '.join(str(c) for r in rows for c in r if c).lower()
        has_interest = any(k in text_all for k in ['出库', 'toc', '订单处理', '附加', '燃油', 'fuel', 'gofo', 'xlm'])
        if has_interest or len(rows) < 100:
            print('\n--- Sheet: %s (%d rows) ---' % (sn, len(rows)))
            for ri, row in enumerate(rows[:15]):
                cells = [str(c)[:20] if c else '' for c in row]
                non_empty = [c for c in cells if c.strip()]
                if non_empty:
                    print('  R%d: %s' % (ri, ' | '.join(cells[:10])))
    wb.close()

for f in leecang_new:
    print('\n\n=== LEECANG NEW: %s ===' % os.path.basename(f))
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for sn in wb.sheetnames:
        if 'FEDEX' in sn.upper() or 'GOFO' in sn.upper() or 'XLM' in sn.upper():
            ws = wb[sn]
            rows = list(ws.iter_rows(values_only=True))
            print('\n--- Sheet: %s (%d rows) ---' % (sn, len(rows)))
            for ri, row in enumerate(rows[:20]):
                cells = [str(c)[:20] if c else '' for c in row]
                non_empty = [c for c in cells if c.strip()]
                if non_empty:
                    print('  R%d: %s' % (ri, ' | '.join(cells[:10])))
    wb.close()

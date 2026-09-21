import sys, os, re
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)

path = r'E:\新建文件夹\海外仓\正式价格'
files = sorted([f for f in os.listdir(path) if f.endswith('.xlsx')])
import openpyxl

# File 1: 乐歌 出库操作 - why outbound=0?
print("=== FILE 1: 乐歌 outbound debug ===")
wb = openpyxl.load_workbook(os.path.join(path, files[0]), read_only=True, data_only=True)
ws = wb["海外仓出入库操作"]
rows = list(ws.iter_rows(values_only=True))
text_all = " ".join(str(c) for row in rows for c in row if c).lower()

# Check keywords
for k in ["订单处理", "出库费", "出库服务", "toc", "一件代发"]:
    print("  Has '%s': %s" % (k, k in text_all))

# Check _is_tier detection
tier_pattern = re.compile(r"\d+.*[<≤].*\d+|W\s*[>≥]\s*\d+|≤\s*\d+", re.I)
for ri, row in enumerate(rows[:25]):
    for ci, c in enumerate(row):
        if c and isinstance(c, str) and tier_pattern.search(c.strip()):
            print("  Tier at R%dC%d: %s" % (ri+1, ci, c.strip()[:40]))
wb.close()

# File 2: SMART outbound debug
print("\n=== FILE 2: SMART outbound debug ===")
wb = openpyxl.load_workbook(os.path.join(path, files[1]), read_only=True, data_only=True)
ws = wb["库内操作"]
rows = list(ws.iter_rows(values_only=True))

# Find first few weight tiers
tier_count = 0
for ri, row in enumerate(rows):
    for ci, c in enumerate(row):
        if c and isinstance(c, str) and tier_pattern.search(c.strip()):
            # Show context
            nearby = [str(row[j])[:20] if j < len(row) and row[j] else "" for j in range(max(0,ci-2), min(ci+3, len(row)))]
            print("  R%dC%d: %s  ctx=%s" % (ri+1, ci, c.strip()[:30], nearby))
            tier_count += 1
            if tier_count >= 10: break
    if tier_count >= 10: break

# Find the actual outbound section
print("\n  Looking for '订单处理' section...")
for ri, row in enumerate(rows):
    text = " ".join(str(c) for c in row if c)
    if "订单处理" in text or "出库费" in text:
        print("  Found at R%d: %s" % (ri+1, text[:80]))
        # Show next 15 rows
        for ri2 in range(ri, min(ri+15, len(rows))):
            r = rows[ri2]
            vals = ["[%d]%s" % (ci, str(c)[:20]) for ci, c in enumerate(r[:12]) if c is not None and str(c).strip()]
            print("    R%d: %s" % (ri2+1, ", ".join(vals)))
        break
wb.close()

# File 4: 实力派 outbound debug
print("\n=== FILE 4: 实力派 outbound debug ===")
wb = openpyxl.load_workbook(os.path.join(path, files[3]), read_only=True, data_only=True)
ws = wb["库内费用"]
rows = list(ws.iter_rows(values_only=True))
# Show rows 7-18 (outbound section)
for ri in range(6, 18):
    if ri < len(rows):
        row = rows[ri]
        vals = ["[%d]%s" % (ci, str(c)[:20]) for ci, c in enumerate(row[:10]) if c is not None and str(c).strip()]
        print("  R%d: %s" % (ri+1, ", ".join(vals)))
wb.close()

# File 5: 安美 outbound debug
print("\n=== FILE 5: 安美 outbound debug ===")
wb = openpyxl.load_workbook(os.path.join(path, files[4]), read_only=True, data_only=True)
ws = wb["一件代发"]
rows = list(ws.iter_rows(values_only=True))
# Show rows 13-27 (outbound section)
for ri in range(12, 27):
    if ri < len(rows):
        row = rows[ri]
        vals = ["[%d]%s" % (ci, str(c)[:25]) for ci, c in enumerate(row[:9]) if c is not None and str(c).strip()]
        print("  R%d: %s" % (ri+1, ", ".join(vals)))
wb.close()

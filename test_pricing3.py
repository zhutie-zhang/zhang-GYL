import sys, os, re
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
import pricing

path = r'E:\新建文件夹\海外仓\正式价格'
files = sorted([f for f in os.listdir(path) if f.endswith('.xlsx')])

# Debug: check what keyword matching happens for each sheet
import openpyxl
for fi, f in enumerate(files):
    print("=== FILE %d: %s ===" % (fi+1, f[:40]))
    wb = openpyxl.load_workbook(os.path.join(path, f), read_only=True, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        text_all = pricing._all_text(rows).lower()
        # Check which keywords match
        matched = []
        if any(k in text_all for k in ["存储", "仓储费", "storage", "仓租"]): matched.append("storage")
        if any(k in text_all for k in ["订单处理", "出库费", "一件代发", "代发", "toc订单", "toc"]): matched.append("outbound")
        if any(k in text_all for k in ["入库", "卸货", "卸柜", "receiving", "devanning"]): matched.append("inbound")
        if any(k in text_all for k in ["增值服务", "value added"]): matched.append("value_added")
        if any(k in text_all for k in ["附加费", "surcharge", "ahs"]): matched.append("surcharge")
        if pricing._looks_like_zone(rows): matched.append("zone_shipping")
        if any(k in text_all for k in ["库内", "操作费"]): matched.append("warehouse_ops")
        # Skip check
        if any(k in text_all for k in ["目录", "简介", "基础信息", "客户信息", "索赔", "清关", "提货", "customs", "trucking", "claimer", "仓储定义"]): matched.append("SKIP")
        print("  Sheet '%s' (rows=%d): %s" % (sn, len(rows), matched))
    wb.close()
    print()

print("=== Testing outbound parse on 乐歌 sheet ===")
wb = openpyxl.load_workbook(os.path.join(path, files[0]), read_only=True, data_only=True)
ws = wb["海外仓出入库操作"]
rows = list(ws.iter_rows(values_only=True))
# Show first 25 rows with col indices
for ri, row in enumerate(rows[:25]):
    vals = ["[%d]%s" % (ci, str(c)[:20]) for ci, c in enumerate(row[:8]) if c is not None and str(c).strip()]
    print("R%d: %s" % (ri+1, ", ".join(vals)))
wb.close()

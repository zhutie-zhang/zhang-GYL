import sys, os, re, glob
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
import openpyxl

# SMART: check 件出库处理费 + surcharge details on FedEx_AHS and GOFO
files = glob.glob(r'E:\新建文件夹\海外仓\正式价格\*.xlsx')
smart = [f for f in files if 'SMART' in os.path.basename(f) and not os.path.basename(f).startswith('~$')]
leecang_new = [f for f in files if '新乐舱' in os.path.basename(f) and not os.path.basename(f).startswith('~$')]

for f in smart:
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    print('=== SMART ===')
    
    # 件出库处理费
    if '件出库处理费' in wb.sheetnames:
        ws = wb['件出库处理费']
        rows = list(ws.iter_rows(values_only=True))
        print('\n--- 件出库处理费 (%d rows) ---' % len(rows))
        for ri, row in enumerate(rows[:20]):
            cells = [str(c)[:20] if c else '' for c in row]
            non_empty = [c for c in cells if c.strip()]
            if non_empty:
                print('  R%d: %s' % (ri, ' | '.join(cells[:10])))

    # FedEx_AHS right-side surcharges
    if 'Fedex_AHS' in wb.sheetnames:
        ws = wb['Fedex_AHS']
        rows = list(ws.iter_rows(values_only=True))
        print('\n--- Fedex_AHS surcharges (right side) ---')
        for ri, row in enumerate(rows[13:30]):
            cells = [str(c)[:22] if c else '' for c in row]
            print('  R%d: %s' % (ri+13, ' | '.join(cells[:12])))

    # GOFO Ground right-side surcharges
    if 'GOFO Ground' in wb.sheetnames:
        ws = wb['GOFO Ground']
        rows = list(ws.iter_rows(values_only=True))
        print('\n--- GOFO Ground surcharges (right side) ---')
        for ri, row in enumerate(rows[:25]):
            cells = [str(c)[:22] if c else '' for c in row]
            print('  R%d: %s' % (ri, ' | '.join(cells[:12])))

    wb.close()

for f in leecang_new:
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    print('\n\n=== LEECANG SVIP ===')
    if 'FEDEX- Ground & Home Delivery' in wb.sheetnames:
        ws = wb['FEDEX- Ground & Home Delivery']
        rows = list(ws.iter_rows(values_only=True))
        print('\n--- FEDEX surcharges (right side) ---')
        for ri, row in enumerate(rows[:25]):
            cells = [str(c)[:22] if c else '' for c in row]
            print('  R%d: %s' % (ri, ' | '.join(cells[:12])))
    wb.close()

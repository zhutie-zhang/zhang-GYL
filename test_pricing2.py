import sys, os
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)

path = r'E:\新建文件夹\海外仓\正式价格'
files = sorted([f for f in os.listdir(path) if f.endswith('.xlsx')])

# File 1: 乐歌 - 找出库sheet的完整内容
import openpyxl
print("=== FILE 1: 乐歌 - 出库操作 sheet ===")
wb = openpyxl.load_workbook(os.path.join(path, files[0]), read_only=True, data_only=True)
ws = wb["海外仓出入库操作"]
for ri, row in enumerate(ws.iter_rows(values_only=True)):
    vals = [str(c)[:25] if c is not None else '' for c in row[:11]]
    print("R%d: %s" % (ri+1, vals))
wb.close()

print("\n=== FILE 4: 实力派 - 库内费用 sheet (全部) ===")
wb = openpyxl.load_workbook(os.path.join(path, files[3]), read_only=True, data_only=True)
ws = wb["库内费用"]
for ri, row in enumerate(ws.iter_rows(values_only=True)):
    vals = [str(c)[:25] if c is not None else '' for c in row[:10]]
    print("R%d: %s" % (ri+1, vals))
wb.close()

print("\n=== FILE 4: 实力派 - 附加费 sheet (前40行) ===")
wb = openpyxl.load_workbook(os.path.join(path, files[3]), read_only=True, data_only=True)
ws = wb["附加费"]
for ri, row in enumerate(ws.iter_rows(max_row=40, values_only=True)):
    vals = [str(c)[:30] if c is not None else '' for c in row[:8]]
    print("R%d: %s" % (ri+1, vals))
wb.close()

print("\n=== FILE 5: 安美 - 一件代发 sheet (前40行) ===")
wb = openpyxl.load_workbook(os.path.join(path, files[4]), read_only=True, data_only=True)
ws = wb["一件代发"]
for ri, row in enumerate(ws.iter_rows(max_row=40, values_only=True)):
    vals = [str(c)[:30] if c is not None else '' for c in row[:9]]
    print("R%d: %s" % (ri+1, vals))
wb.close()

print("\n=== FILE 5: 安美 - 仓储 sheet (前30行) ===")
wb = openpyxl.load_workbook(os.path.join(path, files[4]), read_only=True, data_only=True)
ws = wb["仓储"]
for ri, row in enumerate(ws.iter_rows(max_row=30, values_only=True)):
    vals = [str(c)[:30] if c is not None else '' for c in row[:10]]
    print("R%d: %s" % (ri+1, vals))
wb.close()

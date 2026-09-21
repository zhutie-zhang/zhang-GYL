import sys, os, re
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)

path = r'E:\新建文件夹\海外仓\正式价格'
files = sorted([f for f in os.listdir(path) if f.endswith('.xlsx')])

import openpyxl

def all_text(rows):
    return " ".join(str(c) for row in rows for c in row if c)

# Test keyword matching for key sheets
test_sheets = [
    (0, "海外仓出入库操作"),  # 乐歌 outbound
    (1, "库内操作"),  # SMART mixed
    (3, "库内费用"),  # 实力派 mixed
    (4, "一件代发"),  # 安美 outbound
    (3, "附加费"),  # 实力派 surcharges
]

for fi, sn in test_sheets:
    fp = os.path.join(path, files[fi])
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    if sn not in wb.sheetnames:
        print("SKIP: file %d has no sheet '%s'" % (fi+1, sn))
        continue
    ws = wb[sn]
    rows = list(ws.iter_rows(values_only=True))
    text = all_text(rows).lower()
    
    print("=== FILE %d / Sheet '%s' ===" % (fi+1, sn))
    print("  Has '订单处理': %s" % ("订单处理" in text))
    print("  Has '出库费': %s" % ("出库费" in text))
    print("  Has '一件代发': %s" % ("一件代发" in text))
    print("  Has '代发': %s" % ("代发" in text))
    print("  Has 'toc': %s" % ("toc" in text))
    
    # Test outbound keyword match
    if any(k in text for k in ["订单处理", "出库费", "一件代发", "代发", "toc订单", "toc"]):
        print("  -> MATCHES outbound keyword")
    else:
        print("  -> DOES NOT match outbound keyword")
    
    # Check what column has weight tiers
    for ri, row in enumerate(rows[:20]):
        for ci, c in enumerate(row):
            if c and isinstance(c, str):
                s = c.strip()
                if re.search(r"\d+.*[<≤].*\d+|W\s*[>≥]\s*\d+", s, re.I):
                    print("  Weight tier at R%dC%d: '%s'" % (ri+1, ci, s[:40]))
                    # Check next column for price
                    if ci+1 < len(row):
                        print("    Price at R%dC%d: %s" % (ri+1, ci+1, row[ci+1]))
                    break
    wb.close()
    print()

# Also check 实力派 surcharges sheet structure
print("=== FILE 4 / 附加费 sheet structure ===")
wb = openpyxl.load_workbook(os.path.join(path, files[3]), read_only=True, data_only=True)
ws = wb["附加费"]
rows = list(ws.iter_rows(max_row=25, values_only=True))
for ri, row in enumerate(rows):
    vals = [(ci, str(c)[:30]) for ci, c in enumerate(row[:8]) if c is not None and str(c).strip()]
    if vals:
        print("R%d: %s" % (ri+1, vals))
wb.close()

# Check 实力派 基础运费 sheet (should be shipping matrix)
print("\n=== FILE 4 / 基础运费 sheet (first 10 rows) ===")
wb = openpyxl.load_workbook(os.path.join(path, files[3]), read_only=True, data_only=True)
ws = wb["基础运费"]
rows = list(ws.iter_rows(max_row=10, values_only=True))
for ri, row in enumerate(rows):
    vals = [(ci, str(c)[:20]) for ci, c in enumerate(row[:10]) if c is not None and str(c).strip()]
    if vals:
        print("R%d: %s" % (ri+1, vals))
wb.close()

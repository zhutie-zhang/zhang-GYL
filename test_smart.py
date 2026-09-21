import sys, os, re
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
import openpyxl

path = r'E:\新建文件夹\海外仓\正式价格\YFJ&SMART仓储一件代发报价 Eff.260616(1).xlsx'
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2: continue
    print('=== Sheet: %s (%d rows) ===' % (sn, len(rows)))
    for ri, row in enumerate(rows[:12]):
        cells = [str(c)[:18] if c else '' for c in row]
        non_empty = [c for c in cells if c.strip()]
        if non_empty:
            print('  R%d: %s' % (ri, ' | '.join(cells[:10])))
    print()
wb.close()

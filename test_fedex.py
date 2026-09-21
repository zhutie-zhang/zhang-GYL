import sys, os, re
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
import openpyxl

path = r'E:\新建文件夹\海外仓\正式价格\2026乐歌公共海外仓报价-云帆济.xlsx'
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

# Find FedEx Ground sheet
for sn in wb.sheetnames:
    if 'fedex' in sn.lower():
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        print('=== Sheet: %s (%d rows) ===' % (sn, len(rows)))
        for ri, row in enumerate(rows[:20]):
            cells = [str(c)[:25] if c else '' for c in row]
            print('  R%d: %s' % (ri, ' | '.join(cells)))
        # Find 60lb row
        for ri, row in enumerate(rows):
            for ci, c in enumerate(row):
                if c and isinstance(c, (int, float)) and 59 <= c <= 61:
                    cells = [str(x)[:15] if x else '' for x in row]
                    print('  R%d [w=%s]: %s' % (ri, c, ' | '.join(cells)))
                    break
        # Fuel/surcharge
        for ri, row in enumerate(rows):
            text = ' '.join(str(c) for c in row if c)
            if re.search(r'fuel|燃油|surcharge|AHS|residential|住宅|附加|额外', text, re.I):
                cells = [str(x)[:25] if x else '' for x in row]
                print('  R%d [surch]: %s' % (ri, ' | '.join(cells)))
        break
wb.close()

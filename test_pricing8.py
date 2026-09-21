import sys, os, re
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
import pricing

path = r'E:\新建文件夹\海外仓\正式价格'
files = sorted([f for f in os.listdir(path) if f.endswith('.xlsx')])
import openpyxl

tier_re = re.compile(r"\d+.*[<≤].*\d+|W\s*[>≥]\s*\d+|≤\s*\d+", re.I)

# 安美: why only 1 tier at $90120?
print("=== 安美 一件代发 tier detection ===")
wb = openpyxl.load_workbook(os.path.join(path, files[4]), read_only=True, data_only=True)
ws = wb["一件代发"]
rows = list(ws.iter_rows(values_only=True))
for ri, row in enumerate(rows):
    for ci, c in enumerate(row):
        if c and isinstance(c, str) and tier_re.search(c.strip()):
            nearby = [str(row[j])[:20] if j < len(row) and row[j] else "" for j in range(max(0,ci-1), min(ci+3, len(row)))]
            print("  R%dC%d: %s | %s" % (ri+1, ci, str(c)[:30], " ".join(nearby)))
wb.close()

# 实力派: why 75 tiers?
print("\n=== 实力派 库内费用 tier detection ===")
wb = openpyxl.load_workbook(os.path.join(path, files[3]), read_only=True, data_only=True)
ws = wb["库内费用"]
rows = list(ws.iter_rows(values_only=True))
tier_count = 0
for ri, row in enumerate(rows):
    for ci, c in enumerate(row):
        if c and isinstance(c, str) and tier_re.search(c.strip()):
            tier_count += 1
            if tier_count <= 10:
                nearby = [str(row[j])[:20] if j < len(row) and row[j] else "" for j in range(max(0,ci-1), min(ci+3, len(row)))]
                print("  R%dC%d: %s | %s" % (ri+1, ci, str(c)[:30], " ".join(nearby)))
print("  Total tiers found:", tier_count)
wb.close()

# SMART: debug
print("\n=== SMART 库内操作 first tiers ===")
wb = openpyxl.load_workbook(os.path.join(path, files[1]), read_only=True, data_only=True)
ws = wb["库内操作"]
rows = list(ws.iter_rows(values_only=True))
tier_count = 0
for ri, row in enumerate(rows[:50]):
    for ci, c in enumerate(row):
        if c and isinstance(c, str) and tier_re.search(c.strip()):
            tier_count += 1
            if tier_count <= 5:
                nearby = [str(row[j])[:20] if j < len(row) and row[j] else "" for j in range(max(0,ci-1), min(ci+3, len(row)))]
                print("  R%dC%d: %s | %s" % (ri+1, ci, str(c)[:30], " ".join(nearby)))
# Also find where outbound starts
for ri, row in enumerate(rows):
    text = " ".join(str(c) for c in row if c)
    if "TOC下架出库" in text or "订单处理" in text or "出库费" in text:
        print("  Found outbound at R%d: %s" % (ri+1, text[:80]))
        for ri2 in range(ri, min(ri+15, len(rows))):
            r = rows[ri2]
            vals = ["[%d]%s" % (ci, str(c)[:20]) for ci, c in enumerate(r[:8]) if c is not None and str(c).strip()]
            print("    R%d: %s" % (ri2+1, ", ".join(vals)))
        break
wb.close()

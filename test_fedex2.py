import sys, os, re
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', buffering=1)
sys.path.insert(0, r'C:\Users\1\Documents\Default Project')
import openpyxl

path = r'E:\新建文件夹\海外仓\正式价格\2026乐歌公共海外仓报价-云帆济.xlsx'
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

for sn in wb.sheetnames:
    if 'fedex' in sn.lower():
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        print('=== %s ===' % sn)
        # Print all rows that have surcharge data on the right side (col >= 10)
        for ri, row in enumerate(rows):
            right_cells = []
            for ci in range(10, min(len(row), 25)):
                c = row[ci] if ci < len(row) else None
                if c:
                    right_cells.append('C%d=%s' % (ci, str(c)[:40]))
            if right_cells:
                print('  R%d: %s' % (ri, ' | '.join(right_cells)))
        break
wb.close()

#!/usr/bin/env python3
"""
海外仓对账报告生成器
将对账结果生成格式化的 Excel 报告，包含差异明细和费用统计

Usage:
    python3 generate_report.py <report.json> [--output <output.xlsx>]
"""

import json
import sys
import os
import argparse
from datetime import datetime


def generate_excel_report(report_data, output_path):
    """生成 Excel 对账差异报告"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    wb = Workbook()

    # 样式定义
    header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    title_font = Font(name='Arial', bold=True, size=14)
    subtitle_font = Font(name='Arial', bold=True, size=11, color='2F5496')
    normal_font = Font(name='Arial', size=10)
    high_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    medium_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    low_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    summary_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    wrap_align = Alignment(wrap_text=True, vertical='top')

    # ========== Sheet 1: 对账摘要 ==========
    ws1 = wb.active
    ws1.title = '对账摘要'
    summary = report_data.get('summary', {})

    ws1['A1'] = '海外仓账单对账报告'
    ws1['A1'].font = title_font
    ws1.merge_cells('A1:D1')

    ws1['A3'] = '对账日期'
    ws1['B3'] = report_data.get('reconciliation_date', '')
    ws1['A4'] = '账单来源'
    ws1['B4'] = report_data.get('bill_meta', {}).get('source_file', '')
    ws1['A5'] = '账单格式'
    ws1['B5'] = report_data.get('bill_meta', {}).get('format', '')
    ws1['A6'] = '账单行数'
    ws1['B6'] = summary.get('total_items', 0)

    for r in range(3, 7):
        ws1[f'A{r}'].font = subtitle_font

    # 费用金额摘要
    ws1['A8'] = '费用总计'
    ws1['B8'] = summary.get('total_bill_amount', 0)
    ws1['B8'].number_format = '#,##0.00'
    ws1['B8'].fill = summary_fill

    ws1['A9'] = '发现问题数'
    ws1['B9'] = summary.get('total_issues', 0)
    ws1['A10'] = '  高严重度'
    ws1['B10'] = summary.get('high_severity_issues', 0)
    ws1['B10'].fill = high_fill
    ws1['A11'] = '  中严重度'
    ws1['B11'] = summary.get('medium_severity_issues', 0)
    ws1['B11'].fill = medium_fill
    ws1['A12'] = '  低严重度'
    ws1['B12'] = summary.get('low_severity_issues', 0)
    ws1['B12'].fill = low_fill

    ws1['A13'] = '预计多收金额'
    ws1['B13'] = summary.get('estimated_overcharge', 0)
    ws1['B13'].number_format = '#,##0.00'
    ws1['B13'].fill = high_fill
    ws1['B13'].font = Font(name='Arial', bold=True, size=11, color='FF0000')

    # 费用类型分布
    ws1['A15'] = '费用类型分布'
    ws1['A15'].font = subtitle_font

    type_breakdown = summary.get('fee_type_breakdown', {})
    row = 16
    ws1[f'A{row}'] = '费用类型'
    ws1[f'B{row}'] = '金额'
    ws1[f'C{row}'] = '占比'
    for col in ['A', 'B', 'C']:
        ws1[f'{col}{row}'].font = header_font
        ws1[f'{col}{row}'].fill = header_fill
        ws1[f'{col}{row}'].border = thin_border

    total = summary.get('total_bill_amount', 0) or 1
    row = 17
    for fee_type, amount in type_breakdown.items():
        ws1[f'A{row}'] = fee_type
        ws1[f'B{row}'] = amount
        ws1[f'B{row}'].number_format = '#,##0.00'
        ws1[f'C{row}'] = f'{amount/total*100:.1f}%'
        for col in ['A', 'B', 'C']:
            ws1[f'{col}{row}'].font = normal_font
            ws1[f'{col}{row}'].border = thin_border
        row += 1

    # 设置列宽
    ws1.column_dimensions['A'].width = 20
    ws1.column_dimensions['B'].width = 18
    ws1.column_dimensions['C'].width = 12

    # ========== Sheet 2: 差异明细 ==========
    ws2 = wb.create_sheet('差异明细')
    issues = report_data.get('issues', [])

    ws2['A1'] = '对账差异明细'
    ws2['A1'].font = title_font
    ws2.merge_cells('A1:H1')

    headers = ['序号', '规则', '规则名称', '严重度', '账单行', '费用名称', '差异金额', '详细描述']
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=3, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')

    row = 4
    for idx, issue in enumerate(issues, 1):
        severity = issue.get('severity', '')
        fill = high_fill if severity == 'high' else (medium_fill if severity == 'medium' else low_fill)

        amount = issue.get('amount_diff') or issue.get('amount') or 0
        try:
            amount = float(amount)
        except (ValueError, TypeError):
            amount = 0

        values = [
            idx,
            issue.get('rule', ''),
            issue.get('rule_name', ''),
            {'high': '高', 'medium': '中', 'low': '低'}.get(severity, ''),
            issue.get('bill_row', ''),
            issue.get('fee_name', ''),
            amount if amount else '',
            issue.get('description', ''),
        ]

        for col, val in enumerate(values, 1):
            cell = ws2.cell(row=row, column=col, value=val)
            cell.font = normal_font
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = wrap_align
            if col == 7 and isinstance(val, (int, float)) and val:
                cell.number_format = '#,##0.00'

        row += 1

    if not issues:
        ws2.cell(row=4, column=1, value='未发现差异')
        ws2.cell(row=4, column=1).font = normal_font

    # 列宽
    col_widths = [6, 8, 18, 8, 8, 20, 14, 50]
    for i, w in enumerate(col_widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ========== Sheet 3: 原始账单 ==========
    ws3 = wb.create_sheet('原始账单')
    bill_meta = report_data.get('bill_meta', {})

    ws3['A1'] = '原始账单数据'
    ws3['A1'].font = title_font

    ws3['A3'] = '来源文件'
    ws3['B3'] = bill_meta.get('source_file', '')
    ws3['A4'] = '格式'
    ws3['B4'] = bill_meta.get('format', '')
    ws3['A5'] = '总行数'
    ws3['B5'] = bill_meta.get('total_rows', 0)

    # 保存
    wb.save(output_path)
    print(f"Report saved to {output_path}")


def generate_text_summary(report_data):
    """生成文本格式摘要（用于邮件正文）"""
    summary = report_data.get('summary', {})
    issues = report_data.get('issues', {})

    lines = []
    lines.append("=" * 60)
    lines.append("          海外仓账单对账报告")
    lines.append("=" * 60)
    lines.append(f"对账日期: {report_data.get('reconciliation_date', '')}")
    lines.append(f"账单来源: {report_data.get('bill_meta', {}).get('source_file', '')}")
    lines.append("")

    lines.append("【费用统计】")
    lines.append(f"  账单总金额: {summary.get('total_bill_amount', 0):.2f}")
    lines.append(f"  账单总行数: {summary.get('total_items', 0)}")
    lines.append("")

    lines.append("【费用类型分布】")
    for ft, amt in summary.get('fee_type_breakdown', {}).items():
        pct = amt / (summary.get('total_bill_amount', 1) or 1) * 100
        lines.append(f"  {ft}: {amt:.2f} ({pct:.1f}%)")
    lines.append("")

    lines.append("【问题汇总】")
    lines.append(f"  总问题数: {summary.get('total_issues', 0)}")
    lines.append(f"  高严重度: {summary.get('high_severity_issues', 0)}")
    lines.append(f"  中严重度: {summary.get('medium_severity_issues', 0)}")
    lines.append(f"  低严重度: {summary.get('low_severity_issues', 0)}")
    lines.append(f"  预计多收: {summary.get('estimated_overcharge', 0):.2f}")
    lines.append("")

    if issues:
        lines.append("【差异明细】")
        for i, issue in enumerate(issues, 1):
            sev = {'high': '🔴', 'medium': '🟡', 'low': '🟢'}.get(issue.get('severity', ''), '⚪')
            lines.append(f"  {i}. [{sev} {issue.get('rule', '')}] {issue.get('rule_name', '')}")
            lines.append(f"     {issue.get('description', '')}")
            amt = issue.get('amount_diff') or issue.get('amount') or 0
            if amt:
                lines.append(f"     涉及金额: {float(amt):.2f}")
            lines.append("")

    return '\n'.join(lines)


def main():
    parser = argparse.ArgumentParser(description='海外仓对账报告生成器')
    parser.add_argument('report_file', help='对账报告 JSON 文件')
    parser.add_argument('--output', '-o', help='输出 Excel 路径', default=None)
    parser.add_argument('--text', '-t', action='store_true', help='输出文本格式摘要')
    args = parser.parse_args()

    with open(args.report_file, 'r', encoding='utf-8') as f:
        report_data = json.load(f)

    if args.text:
        print(generate_text_summary(report_data))
    else:
        output_path = args.output or args.report_file.replace('.json', '.xlsx')
        generate_excel_report(report_data, output_path)


if __name__ == '__main__':
    main()
